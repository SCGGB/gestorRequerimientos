"""Deriva la plantilla docxtpl desde el machote institucional (TEST-007 / T-001).

Herramienta de desarrollo (no forma parte del runtime del backend). Genera
`templates/ers_plantilla_v1.docx` a partir de `templates/machote_requerimiento.docx`
sin modificar el original (ADR-004).

Uso:
    python tools/derive_ers_template.py

Requiere: python-docx, lxml.
"""

from __future__ import annotations

import copy
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "templates" / "machote_requerimiento.docx"
TARGET = ROOT / "templates" / "ers_plantilla_v1.docx"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
HEADING_STYLES = {"Heading 1", "Heading 2", "Heading 3"}
LIST_STYLES = {"List Paragraph", "List Number"}
FIXED_HEADINGS = {"FIRMAS", "CONTROL DE VERSIONES"}

COLLECTION_SECTIONS = (
    ("Objetivos Generales", "objetivos"),
    ("Consideraciones", "consideraciones"),
    ("Restricciones", "restricciones"),
    ("Volumetr", "volumetria"),
    ("Riesgos Iniciales", "riesgos"),
)

BENEFIT_SECTIONS = (
    ("4.1 Beneficios Cualitativos", "beneficios_cualitativos"),
    ("4.2 Beneficios Cuantitativos", "beneficios_cuantitativos"),
)


def flatten_sdt(docx_path: Path, tmp_path: Path) -> None:
    """Aplana los controles de contenido (w:sdt) de document.xml y encabezados."""
    with zipfile.ZipFile(docx_path) as zin:
        names = zin.namelist()
        payload = {name: zin.read(name) for name in names}

    for name in ("word/document.xml", "word/header1.xml", "word/header2.xml"):
        if name not in payload:
            continue
        root = etree.fromstring(payload[name])
        for sdt in list(root.iter(f"{{{W}}}sdt")):
            content = sdt.find(f"{{{W}}}sdtContent")
            if content is None:
                continue
            parent = sdt.getparent()
            index = parent.index(sdt)
            parent.remove(sdt)
            for offset, child in enumerate(list(content)):
                parent.insert(index + offset, child)
        payload[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in names:
            zout.writestr(name, payload[name])


def set_text(paragraph: Paragraph, text: str) -> None:
    """Reemplaza el texto del parrafo por un unico run (docxtpl lo exige)."""
    p = paragraph._p
    rpr = None
    for run in p.findall(qn("w:r")):
        found = run.find(qn("w:rPr"))
        if found is not None:
            rpr = found
            break
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    run = OxmlElement("w:r")
    if rpr is not None:
        run.append(copy.deepcopy(rpr))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    p.append(run)


def insert_before(paragraph: Paragraph, text: str) -> Paragraph:
    return paragraph.insert_paragraph_before(text)


def insert_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    new_paragraph.add_run(text)
    return new_paragraph


def insert_sequence_after(paragraph: Paragraph, texts: tuple[str, ...]) -> Paragraph:
    anchor = paragraph
    for text in texts:
        anchor = insert_after(anchor, text)
    return anchor


def remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._p.getparent().remove(paragraph._p)


def is_heading(paragraph: Paragraph) -> bool:
    style = paragraph.style.name if paragraph.style is not None else ""
    return style in HEADING_STYLES or paragraph.text.strip() in FIXED_HEADINGS


def find_paragraph(doc: Document, prefix: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise LookupError(f"No se encontro el parrafo que inicia con: {prefix!r}")


def content_run(doc: Document, heading: Paragraph) -> list[Paragraph]:
    """Parrafos de contenido entre un titulo y el siguiente titulo."""
    paragraphs = doc.paragraphs
    start = None
    for index, paragraph in enumerate(paragraphs):
        if paragraph._p is heading._p:
            start = index
            break
    if start is None:
        raise LookupError("Titulo no pertenece al documento")
    result: list[Paragraph] = []
    for paragraph in paragraphs[start + 1:]:
        if is_heading(paragraph):
            break
        result.append(paragraph)
    return result


def pick_body(run: list[Paragraph]) -> Paragraph:
    for paragraph in run:
        if paragraph.text.strip():
            return paragraph
    for paragraph in run:
        if paragraph.style.name in LIST_STYLES:
            return paragraph
    return run[0]


def fill_collection(doc: Document, prefix: str, varname: str) -> None:
    heading = find_paragraph(doc, prefix)
    run = content_run(doc, heading)
    if not run:
        raise LookupError(f"La seccion {prefix!r} no tiene parrafos de contenido")
    body = pick_body(run)
    set_text(body, "{{ item }}")
    for extra in run:
        if extra._p is not body._p:
            remove_paragraph(extra)
    insert_sequence_after(body, ("{%p endfor %}", "{%p else %}", "{{ no_aplica }}", "{%p endif %}"))
    insert_before(body, f"{{%p if {varname} %}}")
    insert_before(body, f"{{%p for item in {varname} %}}")


def fill_paragraph_section(doc: Document, prefix: str, varname: str) -> None:
    heading = find_paragraph(doc, prefix)
    run = content_run(doc, heading)
    if not run:
        raise LookupError(f"La seccion {prefix!r} no tiene parrafos de contenido")
    body = run[0]
    set_text(body, f"{{{{ {varname} }}}}")
    for extra in run[1:]:
        remove_paragraph(extra)


def fill_contextual(doc: Document) -> None:
    heading = find_paragraph(doc, "Diagramas Contextuales")
    run = content_run(doc, heading)
    label_index = None
    for index, paragraph in enumerate(run):
        if paragraph.text.strip().lower().startswith("manejo de errores"):
            label_index = index
            break
    if label_index is None:
        raise LookupError("No se encontro el label 'Manejo de errores'")

    description = None
    for paragraph in run[:label_index]:
        if paragraph.text.strip():
            description = paragraph
            break
    if description is None:
        raise LookupError("No se encontro la descripcion contextual")
    set_text(description, "{{ contexto.descripcion }}")

    label = run[label_index]
    manejo_text = run[label_index + 1] if label_index + 1 < len(run) else None
    if manejo_text is None:
        raise LookupError("No se encontro el texto de manejo de errores")
    set_text(manejo_text, "{{ contexto.manejo_errores_texto }}")
    insert_before(label, "{%p if contexto.manejo_errores %}")
    insert_after(manejo_text, "{%p endif %}")

    for paragraph in run:
        if paragraph._p in (description._p, label._p, manejo_text._p):
            continue
        remove_paragraph(paragraph)


def fill_future_requirements(doc: Document) -> None:
    heading = find_paragraph(doc, "Requisitos futuros")
    run = content_run(doc, heading)
    if not run:
        raise LookupError("La seccion de requisitos futuros esta vacia")
    body = pick_body(run)
    set_text(body, "{{ item }}")
    for extra in run:
        if extra._p is not body._p:
            remove_paragraph(extra)
    insert_sequence_after(
        body,
        ("{%p endfor %}", "{%p else %}", "{{ requisitos_futuros_no_aplica }}", "{%p endif %}"),
    )
    insert_before(body, "{%p if requisitos_futuros %}")
    insert_before(body, "{%p for item in requisitos_futuros %}")


def fill_version_control_table(doc: Document) -> None:
    table = doc.tables[0]
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    for_tag_row = table.add_row()
    set_text(for_tag_row.cells[0].paragraphs[0], "{%tr for v in versiones %}")
    data_row = table.add_row()
    for cell, tag in zip(
        data_row.cells,
        ("{{ v.version }}", "{{ v.fecha }}", "{{ v.descripcion_cambio }}", "{{ v.reviso }}", "{{ v.elaboro }}"),
    ):
        set_text(cell.paragraphs[0], tag)
    table.add_row()
    set_text(table.rows[-1].cells[0].paragraphs[0], "{%tr endfor %}")


def fill_signature_table(doc: Document) -> None:
    table = doc.tables[1]
    rows = list(table.rows)
    names = ("{{ firmas.elaboro.nombre }}", "{{ firmas.reviso.nombre }}", "{{ firmas.autorizo.nombre }}")
    positions = ("{{ firmas.elaboro.cargo }}", "{{ firmas.reviso.cargo }}", "{{ firmas.autorizo.cargo }}")
    for cell, tag in zip(rows[1].cells, names):
        set_text(cell.paragraphs[0], tag)
    for cell, tag in zip(rows[-1].cells, positions):
        set_text(cell.paragraphs[0], tag)


def tag_headers(doc: Document) -> None:
    for section in doc.sections:
        for header in (section.header, section.first_page_header):
            for table in header.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            text = paragraph.text
                            if "VERSIÓN DOCUMENTO" in text:
                                set_text(paragraph, "VERSIÓN DOCUMENTO:  {{ document_version }}")
                            elif "T#" in text:
                                set_text(paragraph, "{{ ticket }}")


def clean_empty_body_paragraphs(doc: Document) -> None:
    for paragraph in list(doc.paragraphs):
        if paragraph.text.strip():
            continue
        if is_heading(paragraph):
            continue
        if paragraph.style.name in LIST_STYLES:
            remove_paragraph(paragraph)


def derive() -> None:
    tmp = TARGET.with_suffix(".flattened.docx")
    flatten_sdt(SOURCE, tmp)
    doc = Document(str(tmp))

    fill_version_control_table(doc)
    fill_signature_table(doc)

    fill_paragraph_section(doc, "Necesidad Actual", "necesidad_actual")
    for prefix, varname in COLLECTION_SECTIONS:
        fill_collection(doc, prefix, varname)
    for prefix, varname in BENEFIT_SECTIONS:
        fill_collection(doc, prefix, varname)
    fill_contextual(doc)
    fill_paragraph_section(doc, "Usuarios y Permisos", "usuarios_permisos")
    fill_future_requirements(doc)

    clean_empty_body_paragraphs(doc)
    tag_headers(doc)

    doc.save(str(TARGET))
    tmp.unlink(missing_ok=True)
    print(f"Plantilla generada: {TARGET}")


if __name__ == "__main__":
    derive()
